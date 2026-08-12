"""
rounds_excel.py — Quantitative round log workbook.

Updated automatically when a round is archived (prepare_fresh_round / archive).
Workbook: reports/kalshi_rounds.xlsx

Sheets:
  Rounds      — one row per archived session (master scorecard)
  By_Asset    — per-round asset breakdown
  Config      — profile / risk fingerprint
  Entry_Bands — win/loss stats by entry-price band
  Legend      — column definitions

CLI:
  python -m kalshi_bot.rounds_excel --rebuild
  python -m kalshi_bot.rounds_excel --session sessions/session_YYYY-MM-DD_HHMM
"""

from __future__ import annotations

import argparse
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("kalshi_bot.rounds_excel")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SESSIONS_DIR = PROJECT_ROOT / "sessions"
REPORTS_DIR = PROJECT_ROOT / "reports"
DEFAULT_XLSX = REPORTS_DIR / "kalshi_rounds.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
RED_FILL = PatternFill("solid", fgColor="FFDCD4")
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

ROUNDS_HEADERS = [
    "session_tag",
    "archive_dir",
    "archived_at",
    "mode",
    "profile",
    "started_at",
    "first_trade_ts",
    "last_trade_ts",
    "duration_hours",
    "start_$",
    "trading_end_$",
    "vault_end_$",
    "equity_end_$",
    "peak_trading_$",
    "peak_equity_$",
    "pnl_equity_$",
    "pnl_equity_%",
    "pnl_trading_only_$",
    "vault_skim_count",
    "vault_skim_total_$",
    "trades",
    "wins",
    "losses",
    "win_rate_%",
    "sharpe_premium",
    "expectancy_$",
    "median_pnl_$",
    "avg_win_$",
    "avg_loss_$",
    "payoff_ratio",
    "profit_factor",
    "best_trade_$",
    "worst_trade_$",
    "mean_ret_premium",
    "std_ret_premium",
    "pct_full_premium_loss",
    "max_equity_dd_%",
    "max_trading_dd_%",
    "max_consec_losses",
    "pct_entry_in_band_15_85",
    "pct_entry_lt_0.15",
    "pct_entry_lt_0.20",
    "avg_entry_wins",
    "avg_entry_losses",
    "btc_trades",
    "btc_pnl_$",
    "btc_wr_%",
    "eth_trades",
    "eth_pnl_$",
    "eth_wr_%",
    "sol_trades",
    "sol_pnl_$",
    "sol_wr_%",
    "xrp_trades",
    "xrp_pnl_$",
    "xrp_wr_%",
    "doge_trades",
    "doge_pnl_$",
    "doge_wr_%",
    "bnb_trades",
    "bnb_pnl_$",
    "bnb_wr_%",
    "hype_trades",
    "hype_pnl_$",
    "hype_wr_%",
    "near_trades",
    "near_pnl_$",
    "near_wr_%",
    "zec_trades",
    "zec_pnl_$",
    "zec_wr_%",
    "has_trades",
]

CONFIG_HEADERS = [
    "session_tag",
    "archive_dir",
    "profile",
    "mode_DRY_RUN",
    "KELLY_FRACTION",
    "MIN_EDGE_PCT",
    "MAX_POS_PCT",
    "PORTFOLIO_GROSS_CAP",
    "MIN_ENTRY_PRICE",
    "MAX_ENTRY_PRICE",
    "MAX_DRAWDOWN_PCT",
    "DRAWDOWN_USE_EQUITY",
    "DRAWDOWN_HALT_ENABLED",
    "MAX_DAILY_LOSS_PCT",
    "MAX_CONSEC_LOSSES",
    "COOLDOWN_MINUTES",
    "ACTIVITY_MANDATE_ENABLED",
    "ACTIVITY_IDLE_SECS",
    "ACTIVITY_PROBE_SIZE_PCT",
    "ALPHA_EDGE_ENABLED",
    "ALPHA_EDGE_MIN",
    "SPOT_CONFIDENCE_MIN",
    "CWM_MIN",
    "LAG_ABSENT_MIN",
    "VOL_HI",
    "EARLY_EXIT_ENABLED",
    "SIM_BALANCE",
    "assets_enabled",
]

ASSET_HEADERS = [
    "session_tag",
    "archive_dir",
    "asset",
    "trades",
    "wins",
    "losses",
    "win_rate_%",
    "pnl_$",
    "avg_pnl_$",
    "best_$",
    "worst_$",
    "avg_entry",
    "share_of_round_pnl_%",
]

BAND_HEADERS = [
    "session_tag",
    "archive_dir",
    "entry_band",
    "trades",
    "wins",
    "losses",
    "win_rate_%",
    "pnl_$",
    "avg_pnl_$",
]


def _load_json(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (json.JSONDecodeError, OSError):
        return {}


def _load_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            out.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return out


def _parse_ts(ts: Optional[str]) -> Optional[datetime]:
    if not ts:
        return None
    try:
        return datetime.fromisoformat(str(ts).replace("Z", "+00:00")).replace(tzinfo=None)
    except (TypeError, ValueError):
        return None


def _safe_float(v: Any, default: float = 0.0) -> float:
    try:
        if v is None:
            return default
        return float(v)
    except (TypeError, ValueError):
        return default


def _entry_band(entry: float) -> str:
    if entry < 0.15:
        return "0.00-0.15"
    if entry < 0.20:
        return "0.15-0.20"
    if entry < 0.40:
        return "0.20-0.40"
    if entry < 0.60:
        return "0.40-0.60"
    if entry < 0.85:
        return "0.60-0.85"
    return "0.85-1.00"


def _sharpe(returns: list[float]) -> float:
    if len(returns) < 2:
        return 0.0
    r = np.array(returns, dtype=float)
    sd = float(np.std(r))
    if sd == 0:
        return 0.0
    return float(np.mean(r) / sd)


def _max_dd(path: list[float]) -> float:
    if not path:
        return 0.0
    peak = path[0]
    mdd = 0.0
    for x in path:
        peak = max(peak, x)
        if peak > 0:
            mdd = max(mdd, (peak - x) / peak)
    return mdd


def _max_consec_losses(pnls: list[float]) -> int:
    best = cur = 0
    for p in pnls:
        if p <= 0:
            cur += 1
            best = max(best, cur)
        else:
            cur = 0
    return best


def summarize_session(archive_dir: Path) -> Optional[dict[str, Any]]:
    """Build a full quantitative summary for one archived session folder."""
    archive_dir = Path(archive_dir).resolve()
    if not archive_dir.is_dir():
        return None

    sim = _load_json(archive_dir / "kalshi_sim.json")
    meta = _load_json(archive_dir / "session_meta.json")
    trades_all = _load_jsonl(archive_dir / "kalshi_trades.jsonl")
    trades = [t for t in trades_all if t.get("side")] or trades_all
    skims = _load_jsonl(archive_dir / "vault_skims.jsonl")

    cfg_levels_raw = meta.get("config_levels_at_archive") or meta.get("config_levels") or {}
    # Older archives sometimes stored a profile string instead of a dict
    if isinstance(cfg_levels_raw, str):
        cfg_levels = {"profile": cfg_levels_raw}
    elif isinstance(cfg_levels_raw, dict):
        cfg_levels = cfg_levels_raw
    else:
        cfg_levels = {}
    summary_meta = meta.get("summary") if isinstance(meta.get("summary"), dict) else {}

    tag = meta.get("session_tag") or archive_dir.name
    start = _safe_float(
        sim.get("starting_balance", summary_meta.get("starting_balance")),
        _safe_float(cfg_levels.get("SIM_BALANCE"), 1000.0),
    )
    trading_end = _safe_float(sim.get("balance"), start)
    vault_end = _safe_float(sim.get("vault_balance"), 0.0)
    if trades and vault_end == 0.0 and trades[-1].get("vault") is not None:
        vault_end = _safe_float(trades[-1].get("vault"), 0.0)
    equity_end = _safe_float(sim.get("total_equity"), trading_end + vault_end)
    peak_trading = _safe_float(sim.get("peak_balance"), trading_end)
    peak_equity = _safe_float(sim.get("peak_equity"), equity_end)

    pnls = [_safe_float(t.get("pnl")) for t in trades]
    wins_p = [p for p in pnls if p > 0]
    losses_p = [p for p in pnls if p <= 0]
    wins = len(wins_p)
    losses = len(losses_p)
    n = len(trades)
    wr = (wins / n * 100.0) if n else 0.0

    returns = []
    for t in trades:
        ref = _safe_float(t.get("entry")) * _safe_float(t.get("contracts"))
        if ref > 0:
            returns.append(_safe_float(t.get("pnl")) / ref)

    # Equity / trading paths
    eq_path = [start]
    tr_path = [start]
    eq = start
    tr = start
    # Reconstruct: prefer logged equity/balance; vault skims reduce trading without changing equity
    for t in trades:
        p = _safe_float(t.get("pnl"))
        if t.get("equity") is not None:
            eq = _safe_float(t.get("equity"))
        else:
            eq += p
        if t.get("balance") is not None:
            tr = _safe_float(t.get("balance"))
        else:
            tr += p
        eq_path.append(eq)
        tr_path.append(tr)
        peak_equity = max(peak_equity, eq)
        peak_trading = max(peak_trading, tr)

    first_ts = trades[0].get("ts") if trades else None
    last_ts = trades[-1].get("ts") if trades else None
    d0, d1 = _parse_ts(first_ts), _parse_ts(last_ts)
    duration_h = round((d1 - d0).total_seconds() / 3600.0, 2) if d0 and d1 else 0.0

    dry = cfg_levels.get("DRY_RUN")
    if dry is None:
        dry = True
    mode = "Paper" if dry else "Live"

    entries = [_safe_float(t.get("entry")) for t in trades]
    in_band = sum(1 for e in entries if 0.15 <= e <= 0.85)
    lt15 = sum(1 for e in entries if e < 0.15)
    lt20 = sum(1 for e in entries if e < 0.20)
    ent_w = [_safe_float(t.get("entry")) for t in trades if _safe_float(t.get("pnl")) > 0]
    ent_l = [_safe_float(t.get("entry")) for t in trades if _safe_float(t.get("pnl")) <= 0]

    avg_win = float(np.mean(wins_p)) if wins_p else 0.0
    avg_loss = float(np.mean(losses_p)) if losses_p else 0.0
    gross_wins = float(sum(wins_p))
    gross_losses = abs(float(sum(losses_p)))
    profit_factor = (gross_wins / gross_losses) if gross_losses > 0 else (999.0 if gross_wins > 0 else 0.0)
    payoff = (avg_win / abs(avg_loss)) if avg_loss < 0 else (999.0 if avg_win > 0 else 0.0)

    # Per asset
    by_asset: dict[str, list[dict]] = {}
    try:
        from .config import all_asset_symbols
        for sym in all_asset_symbols():
            by_asset[sym] = []
    except Exception:
        by_asset = {"BTC": [], "ETH": [], "SOL": [], "XRP": []}
    for t in trades:
        a = str(t.get("asset") or "?")
        by_asset.setdefault(a, []).append(t)

    def asset_stats(rows: list[dict]) -> tuple[int, float, float]:
        if not rows:
            return 0, 0.0, 0.0
        ap = [_safe_float(t.get("pnl")) for t in rows]
        aw = sum(1 for p in ap if p > 0)
        return len(rows), float(sum(ap)), (aw / len(rows) * 100.0)

    def _as(sym: str) -> tuple[int, float, float]:
        return asset_stats(by_asset.get(sym, []))

    btc_n, btc_pnl, btc_wr = _as("BTC")
    eth_n, eth_pnl, eth_wr = _as("ETH")
    sol_n, sol_pnl, sol_wr = _as("SOL")
    xrp_n, xrp_pnl, xrp_wr = _as("XRP")
    doge_n, doge_pnl, doge_wr = _as("DOGE")
    bnb_n, bnb_pnl, bnb_wr = _as("BNB")
    hype_n, hype_pnl, hype_wr = _as("HYPE")
    near_n, near_pnl, near_wr = _as("NEAR")
    zec_n, zec_pnl, zec_wr = _as("ZEC")

    # Entry bands
    bands: dict[str, list[float]] = {}
    for t in trades:
        b = _entry_band(_safe_float(t.get("entry")))
        bands.setdefault(b, []).append(_safe_float(t.get("pnl")))

    band_rows = []
    for band, vals in sorted(bands.items()):
        w = sum(1 for v in vals if v > 0)
        l_ = len(vals) - w
        band_rows.append({
            "session_tag": tag,
            "archive_dir": str(archive_dir.relative_to(PROJECT_ROOT)),
            "entry_band": band,
            "trades": len(vals),
            "wins": w,
            "losses": l_,
            "win_rate_%": round(w / len(vals) * 100.0, 2) if vals else 0.0,
            "pnl_$": round(sum(vals), 4),
            "avg_pnl_$": round(float(np.mean(vals)), 4) if vals else 0.0,
        })

    asset_rows = []
    round_pnl = float(sum(pnls)) if pnls else 0.0
    # Always emit a row per configured asset (0 trades) so new markets show up in Excel
    for asset in sorted(by_asset.keys()):
        rows = by_asset.get(asset) or []
        if not rows:
            asset_rows.append({
                "session_tag": tag,
                "archive_dir": str(archive_dir.relative_to(PROJECT_ROOT)),
                "asset": asset,
                "trades": 0,
                "wins": 0,
                "losses": 0,
                "win_rate_%": 0.0,
                "pnl_$": 0.0,
                "avg_pnl_$": 0.0,
                "best_$": 0.0,
                "worst_$": 0.0,
                "avg_entry": 0.0,
                "share_of_round_pnl_%": 0.0,
            })
            continue
        ap = [_safe_float(t.get("pnl")) for t in rows]
        aw = sum(1 for p in ap if p > 0)
        ae = [_safe_float(t.get("entry")) for t in rows]
        asset_rows.append({
            "session_tag": tag,
            "archive_dir": str(archive_dir.relative_to(PROJECT_ROOT)),
            "asset": asset,
            "trades": len(rows),
            "wins": aw,
            "losses": len(rows) - aw,
            "win_rate_%": round(aw / len(rows) * 100.0, 2),
            "pnl_$": round(float(sum(ap)), 4),
            "avg_pnl_$": round(float(np.mean(ap)), 4),
            "best_$": round(max(ap), 4),
            "worst_$": round(min(ap), 4),
            "avg_entry": round(float(np.mean(ae)), 4) if ae else 0.0,
            "share_of_round_pnl_%": round(float(sum(ap)) / round_pnl * 100.0, 2) if round_pnl else 0.0,
        })

    skim_total = sum(_safe_float(s.get("amount")) for s in skims)

    pnl_equity = equity_end - start
    round_row = {
        "session_tag": tag,
        "archive_dir": str(archive_dir.relative_to(PROJECT_ROOT)),
        "archived_at": meta.get("archived_at") or "",
        "mode": mode,
        "profile": cfg_levels.get("profile") or "",
        "started_at": meta.get("started_at") or "",
        "first_trade_ts": first_ts or "",
        "last_trade_ts": last_ts or "",
        "duration_hours": duration_h,
        "start_$": round(start, 4),
        "trading_end_$": round(trading_end, 4),
        "vault_end_$": round(vault_end, 4),
        "equity_end_$": round(equity_end, 4),
        "peak_trading_$": round(peak_trading, 4),
        "peak_equity_$": round(max(peak_equity, equity_end), 4),
        "pnl_equity_$": round(pnl_equity, 4),
        "pnl_equity_%": round(pnl_equity / start * 100.0, 4) if start else 0.0,
        "pnl_trading_only_$": round(trading_end - start, 4),
        "vault_skim_count": len(skims),
        "vault_skim_total_$": round(skim_total, 4),
        "trades": n,
        "wins": wins,
        "losses": losses,
        "win_rate_%": round(wr, 4),
        "sharpe_premium": round(_sharpe(returns), 4),
        "expectancy_$": round(float(np.mean(pnls)), 4) if pnls else 0.0,
        "median_pnl_$": round(float(np.median(pnls)), 4) if pnls else 0.0,
        "avg_win_$": round(avg_win, 4),
        "avg_loss_$": round(avg_loss, 4),
        "payoff_ratio": round(payoff, 4),
        "profit_factor": round(profit_factor, 4),
        "best_trade_$": round(max(pnls), 4) if pnls else 0.0,
        "worst_trade_$": round(min(pnls), 4) if pnls else 0.0,
        "mean_ret_premium": round(float(np.mean(returns)), 4) if returns else 0.0,
        "std_ret_premium": round(float(np.std(returns)), 4) if returns else 0.0,
        "pct_full_premium_loss": round(
            sum(1 for r in returns if r <= -0.999) / len(returns) * 100.0, 2
        ) if returns else 0.0,
        "max_equity_dd_%": round(_max_dd(eq_path) * 100.0, 4),
        "max_trading_dd_%": round(_max_dd(tr_path) * 100.0, 4),
        "max_consec_losses": _max_consec_losses(pnls),
        "pct_entry_in_band_15_85": round(in_band / n * 100.0, 2) if n else 0.0,
        "pct_entry_lt_0.15": round(lt15 / n * 100.0, 2) if n else 0.0,
        "pct_entry_lt_0.20": round(lt20 / n * 100.0, 2) if n else 0.0,
        "avg_entry_wins": round(float(np.mean(ent_w)), 4) if ent_w else 0.0,
        "avg_entry_losses": round(float(np.mean(ent_l)), 4) if ent_l else 0.0,
        "btc_trades": btc_n,
        "btc_pnl_$": round(btc_pnl, 4),
        "btc_wr_%": round(btc_wr, 2),
        "eth_trades": eth_n,
        "eth_pnl_$": round(eth_pnl, 4),
        "eth_wr_%": round(eth_wr, 2),
        "sol_trades": sol_n,
        "sol_pnl_$": round(sol_pnl, 4),
        "sol_wr_%": round(sol_wr, 2),
        "xrp_trades": xrp_n,
        "xrp_pnl_$": round(xrp_pnl, 4),
        "xrp_wr_%": round(xrp_wr, 2),
        "doge_trades": doge_n,
        "doge_pnl_$": round(doge_pnl, 4),
        "doge_wr_%": round(doge_wr, 2),
        "bnb_trades": bnb_n,
        "bnb_pnl_$": round(bnb_pnl, 4),
        "bnb_wr_%": round(bnb_wr, 2),
        "hype_trades": hype_n,
        "hype_pnl_$": round(hype_pnl, 4),
        "hype_wr_%": round(hype_wr, 2),
        "near_trades": near_n,
        "near_pnl_$": round(near_pnl, 4),
        "near_wr_%": round(near_wr, 2),
        "zec_trades": zec_n,
        "zec_pnl_$": round(zec_pnl, 4),
        "zec_wr_%": round(zec_wr, 2),
        "has_trades": n > 0,
    }

    assets_enabled = cfg_levels.get("assets_enabled")
    if isinstance(assets_enabled, list):
        assets_str = ",".join(str(a) for a in assets_enabled)
    else:
        assets_str = str(assets_enabled or "")

    config_row = {
        "session_tag": tag,
        "archive_dir": str(archive_dir.relative_to(PROJECT_ROOT)),
        "profile": cfg_levels.get("profile") or "",
        "mode_DRY_RUN": dry,
        "KELLY_FRACTION": cfg_levels.get("KELLY_FRACTION"),
        "MIN_EDGE_PCT": cfg_levels.get("MIN_EDGE_PCT"),
        "MAX_POS_PCT": cfg_levels.get("MAX_POS_PCT"),
        "PORTFOLIO_GROSS_CAP": cfg_levels.get("PORTFOLIO_GROSS_CAP"),
        "MIN_ENTRY_PRICE": cfg_levels.get("MIN_ENTRY_PRICE"),
        "MAX_ENTRY_PRICE": cfg_levels.get("MAX_ENTRY_PRICE"),
        "MAX_DRAWDOWN_PCT": cfg_levels.get("MAX_DRAWDOWN_PCT"),
        "DRAWDOWN_USE_EQUITY": cfg_levels.get("DRAWDOWN_USE_EQUITY"),
        "DRAWDOWN_HALT_ENABLED": cfg_levels.get("DRAWDOWN_HALT_ENABLED"),
        "MAX_DAILY_LOSS_PCT": cfg_levels.get("MAX_DAILY_LOSS_PCT"),
        "MAX_CONSEC_LOSSES": cfg_levels.get("MAX_CONSEC_LOSSES"),
        "COOLDOWN_MINUTES": cfg_levels.get("COOLDOWN_MINUTES"),
        "ACTIVITY_MANDATE_ENABLED": cfg_levels.get("ACTIVITY_MANDATE_ENABLED"),
        "ACTIVITY_IDLE_SECS": cfg_levels.get("ACTIVITY_IDLE_SECS"),
        "ACTIVITY_PROBE_SIZE_PCT": cfg_levels.get("ACTIVITY_PROBE_SIZE_PCT"),
        "ALPHA_EDGE_ENABLED": cfg_levels.get("ALPHA_EDGE_ENABLED"),
        "ALPHA_EDGE_MIN": cfg_levels.get("ALPHA_EDGE_MIN"),
        "SPOT_CONFIDENCE_MIN": cfg_levels.get("SPOT_CONFIDENCE_MIN"),
        "CWM_MIN": cfg_levels.get("CWM_MIN"),
        "LAG_ABSENT_MIN": cfg_levels.get("LAG_ABSENT_MIN"),
        "VOL_HI": cfg_levels.get("VOL_HI"),
        "EARLY_EXIT_ENABLED": cfg_levels.get("EARLY_EXIT_ENABLED"),
        "SIM_BALANCE": cfg_levels.get("SIM_BALANCE"),
        "assets_enabled": assets_str,
    }

    return {
        "round": round_row,
        "config": config_row,
        "assets": asset_rows,
        "bands": band_rows,
    }


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def _write_rows(ws, headers: list[str], rows: list[dict]) -> None:
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    _style_header(ws, len(headers))
    for r_i, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = ws.cell(r_i, c, val)
            cell.border = THIN
            cell.alignment = Alignment(horizontal="center")
            if r_i % 2 == 0:
                cell.fill = ALT_FILL
            if h in ("pnl_equity_$", "pnl_$", "btc_pnl_$", "eth_pnl_$", "sol_pnl_$") and isinstance(val, (int, float)):
                if val > 0:
                    cell.fill = GREEN_FILL
                elif val < 0:
                    cell.fill = RED_FILL
        # Highlight equity PnL on Rounds
        if "pnl_equity_$" in headers:
            idx = headers.index("pnl_equity_$") + 1
            v = row.get("pnl_equity_$", 0)
            cell = ws.cell(r_i, idx)
            if isinstance(v, (int, float)):
                cell.fill = GREEN_FILL if v > 0 else (RED_FILL if v < 0 else ALT_FILL)

    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(len(h) + 2, 10), 28)


def _legend_rows() -> list[tuple[str, str]]:
    return [
        ("Rounds sheet", "One row per archived paper/live round — primary scorecard"),
        ("equity_end_$", "Trading balance + vault (true economic result)"),
        ("pnl_equity_$ / %", "Equity − start (use this, not trading-only PnL)"),
        ("pnl_trading_only_$", "Final trading book − start (misleading after vault skims)"),
        ("sharpe_premium", "mean(pnl/premium) / std — bot's built-in Sharpe, not annualized"),
        ("profit_factor", "Gross wins / |gross losses|"),
        ("payoff_ratio", "Avg win / |avg loss|"),
        ("pct_full_premium_loss", "% of trades with ~−100% premium return (binary wipeouts)"),
        ("max_equity_dd_%", "Peak-to-trough drawdown on equity path"),
        ("max_trading_dd_%", "Peak-to-trough on trading balance only"),
        ("pct_entry_lt_0.15", "Lottery-ticket share (Round 12 failure mode)"),
        ("pct_entry_in_band_15_85", "Share inside disciplined entry band"),
        ("By_Asset", "Per-asset contribution — find which coins help/hurt"),
        ("Config", "Fingerprint of risk/gates for that round"),
        ("Entry_Bands", "Where wins/losses cluster by contract price"),
        ("Update trigger", "Written on round archive (stop/log or --fresh-round)"),
    ]


def create_workbook(payloads: list[dict[str, Any]], path: Path = DEFAULT_XLSX) -> Path:
    """Rebuild the full workbook from summarized session payloads."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    # Sort rounds by first trade / archive name
    def sort_key(p: dict) -> str:
        r = p["round"]
        return str(r.get("first_trade_ts") or r.get("archived_at") or r.get("archive_dir") or "")

    payloads = sorted(payloads, key=sort_key)

    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Rounds"
    _write_rows(ws_r, ROUNDS_HEADERS, [p["round"] for p in payloads])

    ws_a = wb.create_sheet("By_Asset")
    asset_rows: list[dict] = []
    for p in payloads:
        asset_rows.extend(p.get("assets") or [])
    _write_rows(ws_a, ASSET_HEADERS, asset_rows)

    ws_c = wb.create_sheet("Config")
    _write_rows(ws_c, CONFIG_HEADERS, [p["config"] for p in payloads])

    ws_b = wb.create_sheet("Entry_Bands")
    band_rows: list[dict] = []
    for p in payloads:
        band_rows.extend(p.get("bands") or [])
    _write_rows(ws_b, BAND_HEADERS, band_rows)

    ws_l = wb.create_sheet("Legend")
    ws_l["A1"] = "Field"
    ws_l["B1"] = "Meaning"
    _style_header(ws_l, 2)
    for i, (k, v) in enumerate(_legend_rows(), 2):
        ws_l.cell(i, 1, k).border = THIN
        ws_l.cell(i, 2, v).border = THIN
    ws_l.column_dimensions["A"].width = 28
    ws_l.column_dimensions["B"].width = 72

    # Comparison helper row counts on Legend
    ws_l.cell(len(_legend_rows()) + 3, 1, "Generated")
    ws_l.cell(len(_legend_rows()) + 3, 2, datetime.now().isoformat(timespec="seconds"))
    ws_l.cell(len(_legend_rows()) + 4, 1, "Rounds logged")
    ws_l.cell(len(_legend_rows()) + 4, 2, len(payloads))

    wb.save(path)
    log.info("Rounds workbook written → %s (%d rounds)", path, len(payloads))
    return path


def iter_session_dirs(sessions_dir: Path = SESSIONS_DIR) -> list[Path]:
    if not sessions_dir.exists():
        return []
    out = []
    for p in sessions_dir.iterdir():
        if not p.is_dir():
            continue
        if (p / "kalshi_trades.jsonl").exists() or (p / "kalshi_sim.json").exists():
            out.append(p)
    return sorted(out, key=lambda x: x.name)


def rebuild_workbook(path: Path = DEFAULT_XLSX, min_trades: int = 0) -> Path:
    """Rescan sessions/ and rebuild the Excel file."""
    payloads = []
    for d in iter_session_dirs():
        try:
            s = summarize_session(d)
        except Exception as e:
            log.warning("Skip %s: %s", d.name, e)
            continue
        if not s:
            continue
        if int(s["round"].get("trades") or 0) < min_trades:
            continue
        payloads.append(s)
    return create_workbook(payloads, path)


def upsert_session(archive_dir: Path, path: Path = DEFAULT_XLSX) -> Path:
    """
    Add/replace one archived session in the workbook.
    Rebuilds from all sessions (simple + consistent) after ensuring archive exists.
    """
    archive_dir = Path(archive_dir)
    if not archive_dir.exists():
        raise FileNotFoundError(archive_dir)
    # Full rebuild keeps sheets consistent (asset/band/config aligned)
    return rebuild_workbook(path=path, min_trades=0)


def record_archived_round(archive_dir: Path) -> Optional[Path]:
    """Called from session_meta after a round is archived. Never raises to callers."""
    try:
        path = upsert_session(archive_dir, DEFAULT_XLSX)
        print(f"Rounds Excel updated: {path}")
        return path
    except PermissionError as e:
        # Excel often locks the main workbook; write a sidecar update file instead.
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        alt = DEFAULT_XLSX.with_name(f"kalshi_rounds_update_{stamp}.xlsx")
        try:
            path = upsert_session(archive_dir, alt)
            log.warning(
                "rounds Excel locked (%s) — wrote sidecar %s", e, path
            )
            print(f"Rounds Excel updated (sidecar): {path}")
            return path
        except Exception as e2:
            log.warning("rounds Excel sidecar update failed: %s", e2)
            return None
    except Exception as e:
        # Windows may raise OSError 13 rather than PermissionError
        err = str(e).lower()
        if "permission denied" in err or "errno 13" in err:
            stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            alt = DEFAULT_XLSX.with_name(f"kalshi_rounds_update_{stamp}.xlsx")
            try:
                path = upsert_session(archive_dir, alt)
                log.warning(
                    "rounds Excel locked (%s) — wrote sidecar %s", e, path
                )
                print(f"Rounds Excel updated (sidecar): {path}")
                return path
            except Exception as e2:
                log.warning("rounds Excel sidecar update failed: %s", e2)
                return None
        log.warning("rounds Excel update failed: %s", e)
        return None


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    ap = argparse.ArgumentParser(description="Kalshi rounds Excel log")
    ap.add_argument("--rebuild", action="store_true", help="Rescan sessions/ into reports/kalshi_rounds.xlsx")
    ap.add_argument("--session", type=str, default="", help="Upsert one session folder")
    ap.add_argument("--out", type=str, default=str(DEFAULT_XLSX), help="Output xlsx path")
    ap.add_argument("--min-trades", type=int, default=0, help="Only include rounds with >= N trades (rebuild filter)")
    args = ap.parse_args()
    out = Path(args.out)

    if args.session:
        p = Path(args.session)
        if not p.is_absolute():
            p = PROJECT_ROOT / p
        path = upsert_session(p, out)
        print(f"Updated {path}")
        return

    path = rebuild_workbook(out, min_trades=args.min_trades)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
